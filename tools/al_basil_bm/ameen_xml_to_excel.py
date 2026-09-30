#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Convert an Al-Ameen Stock Taking XML material file to XLSX.

Usage:
  python ameen_xml_to_excel.py input.xml output.xlsx
"""
from xml.etree import ElementTree as ET
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import sys

ZERO_GUID = "00000000-0000-0000-0000-000000000000"


def t(el, name, default=""):
    node = el.find(name)
    if node is None or node.text is None:
        return default
    return str(node.text).strip()


def n(el, name, default=0):
    raw = t(el, name, "")
    if raw == "":
        return default
    try:
        x = float(raw)
        return int(x) if x.is_integer() else x
    except Exception:
        return default


def convert(src, dst):
    root = ET.parse(src).getroot()
    groups, group_codes = {}, {}
    for g in root.findall('.//G'):
        guid = t(g, 'gPtr')
        if guid:
            groups[guid] = t(g, 'GroupName')
            group_codes[guid] = t(g, 'GroupCode')

    headers = [
        'اسم الصنف','رمز الصنف','اسم المجموعة','رمز المجموعة','الوحدة',
        'الكمية','اخر شراء','رمز الباركود','عامل تحويل الوحدة الثانية',
        'عامل تحويل الوحدة الثالثة'
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = 'لائحة أسعار الأصناف'
    ws.append(headers)

    count = 0
    for m in root.findall('.//M'):
        group_guid = t(m, 'MatGroupGuid', ZERO_GUID)
        ws.append([
            t(m, 'MatName'),
            t(m, 'MatCode'),
            groups.get(group_guid, ''),
            group_codes.get(group_guid, ''),
            t(m, 'MatUnity'),
            n(m, 'SourceQuantity', 0),
            n(m, 'PurchasePrice', 0),
            t(m, 'MatBarCode'),
            n(m, 'MatUnit2Factor', 1),
            n(m, 'MatUnit3Factor', 1),
        ])
        count += 1

    fill = PatternFill('solid', fgColor='0F766E')
    font = Font(color='FFFFFF', bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for i, width in enumerate([34,18,22,16,14,14,14,22,24,24], 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.rightToLeft = True
    wb.save(dst)
    print(f'Converted {count} products -> {dst}')


if __name__ == '__main__':
    if len(sys.argv) != 3:
        raise SystemExit('Usage: python ameen_xml_to_excel.py input.xml output.xlsx')
    convert(sys.argv[1], sys.argv[2])
