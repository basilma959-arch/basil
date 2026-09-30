#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Convert an Excel workbook to the XML format understood by Al-Ameen Stock Taking.

The column "اخر شراء" is the authoritative purchase price. Other price columns are ignored.

Usage:
  python excel_to_ameen_xml.py input.xlsx output.xml
"""
from openpyxl import load_workbook
from xml.etree.ElementTree import Element, SubElement, ElementTree, indent
import sys, uuid

NS_GROUP = uuid.UUID("ec86a5ba-3ff5-4a4d-aa67-14b991c86a30")
NS_PRODUCT = uuid.UUID("093f01e7-7fb2-47f9-9bea-e4ae3c8ab54a")
ZERO_GUID = "00000000-0000-0000-0000-000000000000"

def text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()

def num(v, default=0):
    try:
        if v is None or v == "":
            return default
        return float(v)
    except Exception:
        return default

def convert(src, dst):
    wb = load_workbook(src, read_only=True, data_only=True)
    ws = wb["لائحة أسعار الأصناف"] if "لائحة أسعار الأصناف" in wb.sheetnames else wb[wb.sheetnames[0]]
    headers = [text(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}
    required = ["اسم الصنف", "رمز الصنف", "اسم المجموعة", "الوحدة", "اخر شراء", "رمز الباركود"]
    missing = [x for x in required if x not in idx]
    if missing:
        raise SystemExit("Missing columns: " + ", ".join(missing))

    records, group_first_code = [], {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = text(row[idx["اسم الصنف"]])
        code = text(row[idx["رمز الصنف"]])
        if not name and not code:
            continue
        group = text(row[idx["اسم المجموعة"]])
        gcode = text(row[idx["رمز المجموعة"]]) if "رمز المجموعة" in idx else ""
        if group and gcode and group not in group_first_code:
            group_first_code[group] = gcode
        records.append({
            "name": name,
            "code": code,
            "group": group,
            "unit": text(row[idx["الوحدة"]]),
            "price": num(row[idx["اخر شراء"]]),
            "barcode": text(row[idx["رمز الباركود"]]),
            "f2": num(row[idx["عامل تحويل الوحدة الثانية"]], 1) if "عامل تحويل الوحدة الثانية" in idx else 1,
            "f3": num(row[idx["عامل تحويل الوحدة الثالثة"]], 1) if "عامل تحويل الوحدة الثالثة" in idx else 1,
            "qty": num(row[idx["الكمية"]], 0) if "الكمية" in idx else 0,
        })

    root = Element("AmeenStockImport")
    SubElement(root, "V").text = "3.1"

    groups = sorted({r["group"] for r in records if r["group"]})
    for n, group in enumerate(groups, 1):
        g = SubElement(root, "G")
        SubElement(g, "gPtr").text = str(uuid.uuid5(NS_GROUP, group))
        SubElement(g, "GroupCode").text = group_first_code.get(group) or f"G{n:04d}"
        SubElement(g, "GroupName").text = group
        SubElement(g, "GroupLatinName").text = ""
        SubElement(g, "ParentGuid").text = ZERO_GUID
        SubElement(g, "GroupType").text = "0"

    for r in records:
        m = SubElement(root, "M")
        key = r["code"] or r["barcode"] or r["name"]
        SubElement(m, "mptr").text = str(uuid.uuid5(NS_PRODUCT, key))
        SubElement(m, "MatGroupGuid").text = str(uuid.uuid5(NS_GROUP, r["group"])) if r["group"] else ZERO_GUID
        SubElement(m, "MatCode").text = r["code"]
        SubElement(m, "MatName").text = r["name"]
        SubElement(m, "MatLatinName").text = ""
        SubElement(m, "MatDefUnit").text = "1"
        SubElement(m, "MatBarCode").text = r["barcode"]
        SubElement(m, "MatBarCode2").text = ""
        SubElement(m, "MatBarCode3").text = ""
        SubElement(m, "MatUnity").text = r["unit"]
        SubElement(m, "MatUnit2").text = ""
        SubElement(m, "MatUnit3").text = ""
        SubElement(m, "MatUnit2Factor").text = str(r["f2"])
        SubElement(m, "MatUnit3Factor").text = str(r["f3"])
        SubElement(m, "PurchasePrice").text = str(r["price"])
        SubElement(m, "SourceQuantity").text = str(r["qty"])

    indent(root, space="  ")
    ElementTree(root).write(dst, encoding="utf-8", xml_declaration=True)
    print(f"Converted {len(records)} products / {len(groups)} groups -> {dst}")

if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: python excel_to_ameen_xml.py input.xlsx output.xml")
    convert(sys.argv[1], sys.argv[2])
